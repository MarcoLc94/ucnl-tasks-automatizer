from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..images import fetch_photo
from .charts import render_chart_png
from .common import CoverInfo, reference_entries, split_document
from .markdown import inline_runs, parse, plain

PRIMARY = RGBColor(0x1F, 0x3A, 0x5F)


def render_docx(
    markdown_text: str, cover: CoverInfo, path: Path, work_dir: Path, body_image: Path | None = None,
) -> Path:
    """
    Portada → Introducción → Contenido → Conclusiones → Referencias, cada parte en hoja nueva.
    body_image: si se da (infografía), el contenido es esa imagen a página completa.
    """
    doc = Document()
    _setup_styles(doc)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)
    _add_page_number_footer(doc)

    _add_cover(doc, cover)
    parts = split_document(parse(markdown_text))
    counter = [0]

    if parts["intro"]:
        doc.add_heading("Introducción", level=1)
        _add_blocks(doc, parts["intro"], work_dir, counter)
        _page_break(doc)

    if body_image:
        pic_par = doc.add_paragraph()
        pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_par.add_run().add_picture(str(body_image), height=Inches(8.6))
    else:
        _add_blocks(doc, parts["body"], work_dir, counter)

    if parts["conclusion"]:
        _page_break(doc)
        doc.add_heading("Conclusiones", level=1)
        _add_blocks(doc, parts["conclusion"], work_dir, counter)

    if parts["references"]:
        _page_break(doc)
        doc.add_heading("Referencias", level=1)
        for entry in sorted(reference_entries(parts["references"]), key=lambda t: plain(t).lower()):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)  # sangría francesa (APA)
            _add_runs(p, entry)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_blocks(doc, blocks: list[dict], work_dir: Path, counter: list[int]):
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            # ## de sección → Título 1; ### → Título 2
            doc.add_heading(block["text"].replace("**", ""), level=max(1, block["level"] - 1))
        elif kind == "paragraph":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(p, block["text"])
        elif kind == "callout":
            p = doc.add_paragraph(style="Intense Quote")
            _add_runs(p, block["text"])
        elif kind == "list":
            style = "List Number" if block["ordered"] else "List Bullet"
            for item in block["items"]:
                _add_runs(doc.add_paragraph(style=style), item)
        elif kind == "table":
            _add_table(doc, block["header"], block["rows"])
        elif kind == "image":
            photo = fetch_photo(block["query"], work_dir)
            if photo:
                doc.add_picture(str(photo["path"]), width=Inches(5.8))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(photo["credit"], style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "chart":
            counter[0] += 1
            png = render_chart_png(block["spec"], work_dir / f"chart_{counter[0]}.png")
            if png:
                doc.add_picture(str(png), width=Inches(5.8))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                title = block["spec"].get("titulo") or block["spec"].get("title") or ""
                cap = doc.add_paragraph(f"Figura {counter[0]}. {title}".strip(". "), style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = PRIMARY
        st.font.bold = True


def _add_cover(doc, cover: CoverInfo):
    def centered(text, size, bold=False, space_before=0, color=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        return p

    centered(cover.university, 18, bold=True, space_before=12, color=PRIMARY)
    centered(cover.title, 22, bold=True, space_before=90, color=PRIMARY)

    first = True
    for label, value in cover.fields():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(110 if first else 0)
        p.paragraph_format.space_after = Pt(6)
        r1 = p.add_run(f"{label}: ")
        r1.bold = True
        r1.font.size = Pt(12)
        r2 = p.add_run(value)
        r2.font.size = Pt(12)
        first = False
    _page_break(doc)


def _add_runs(paragraph, text: str):
    for chunk, bold, italic in inline_runs(text):
        r = paragraph.add_run(chunk)
        r.bold = bold or None
        r.italic = italic or None


def _add_table(doc, header: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(header))
    try:
        table.style = doc.styles["Light Grid Accent 1"]
    except KeyError:
        table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text.replace("**", ""))
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            _add_runs(cell.paragraphs[0], text)
    doc.add_paragraph()


def _add_page_number_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    # La portada no lleva número de página
    doc.sections[0].different_first_page_header_footer = True
