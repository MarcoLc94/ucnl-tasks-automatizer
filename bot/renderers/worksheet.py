"""
Actividades que se entregan contestando un Word adjunto por el profesor.

1. outline(): convierte el Word en texto con marcas de ubicación para la IA:
       [P7] ________________
       [T0 R0 C2] (vacía)
2. La IA responde con líneas en el mismo formato (editables en el panel):
       [P7] My name is Linda and I am from Monterrey.
       [T0 R0 C2] She
       [P43] SUBRAYAR: I am shocked by Alice’s opinion.
       [DESPUÉS P5] Texto que va en un párrafo nuevo después de P5
   Las líneas que no empiezan con "[" continúan la respuesta anterior (saltos de línea).
3. fill(): escribe las respuestas en una copia del Word original, conservando su formato.
"""
import copy
import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

_BLANK = re.compile(r"_{3,}")
_LINE = re.compile(r"^\[(DESPU[EÉ]S\s+)?(P\d+|T\d+\s+R\d+\s+C\d+)\]\s*(.*)$", re.I)
_UNDERLINE = re.compile(r"^(SUBRAYAR|SELECCIONAR|MARCAR|CIRCULAR)\s*:\s*", re.I)


def outline(path: Path, max_chars: int = 12000) -> str:
    """Texto del Word con marcas [P#] y [T# R# C#] para que la IA sepa dónde contestar."""
    doc = Document(str(path))
    lines = []
    # Mantener las listas vivas: así cada elemento XML conserva su identidad al comparar
    paragraphs, tables = list(doc.paragraphs), list(doc.tables)
    p_index = {p._p: n for n, p in enumerate(paragraphs)}
    t_index = {t._tbl: n for n, t in enumerate(tables)}
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            n = p_index.get(child)
            text = paragraphs[n].text.strip() if n is not None else ""
            if text:
                lines.append(f"[P{n}] {text}")
        elif tag == "tbl":
            t = t_index.get(child)
            if t is None:
                continue
            table = tables[t]
            lines.append(f"[T{t}] Tabla de {len(table.rows)} filas x {len(table.columns)} columnas:")
            for r, row in enumerate(table.rows):
                seen = set()
                cells = []
                for c, cell in enumerate(row.cells):
                    if id(cell._tc) in seen:  # celdas combinadas
                        continue
                    seen.add(id(cell._tc))
                    cells.append(f"[T{t} R{r} C{c}] {cell.text.strip() or '(vacía)'}")
                lines.append("   " + " | ".join(cells))
    text = "\n".join(lines)
    return text if len(text) <= max_chars else text[:max_chars] + "\n[… documento recortado …]"


def parse_answers(text: str) -> list[dict]:
    answers: list[dict] = []
    for raw in (text or "").splitlines():
        line = raw.rstrip()
        m = _LINE.match(line.strip())
        if m:
            after, loc, value = bool(m.group(1)), m.group(2).upper(), m.group(3).strip()
            item = {"after": after, "value": value}
            if loc.startswith("P"):
                item["p"] = int(loc[1:])
            else:
                t, r, c = map(int, re.findall(r"\d+", loc))
                item.update(t=t, r=r, c=c)
            answers.append(item)
        elif answers and line.strip():
            answers[-1]["value"] += "\n" + line.strip()
    return answers


def _first_run_format(p: Paragraph):
    for r in p.runs:
        if r.text.strip():
            return r
    return p.runs[0] if p.runs else None


def _rebuild(p: Paragraph, segments: list[tuple[str, dict]]):
    """Reemplaza el texto del párrafo por segmentos (texto, formato extra) conservando la fuente del original."""
    base = _first_run_format(p)
    base_rpr = copy.deepcopy(base._r.rPr) if base is not None and base._r.rPr is not None else None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for text, fmt in segments:
        if not text:
            continue
        for i, chunk in enumerate(text.split("\n")):
            if i:
                p.add_run().add_break()
            run = p.add_run(chunk)
            if base_rpr is not None:
                run._r.insert(0, copy.deepcopy(base_rpr))
            if fmt.get("underline"):
                run.underline = True
            if fmt.get("bold"):
                run.bold = True
            if fmt.get("no_underline"):
                run.underline = False


def _norm_quotes(s: str) -> str:
    return s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')


def _fill_paragraph(p: Paragraph, value: str) -> bool:
    text = p.text
    um = _UNDERLINE.match(value)
    if um:
        option = value[um.end():].strip().strip('"')
        idx = _norm_quotes(text).find(_norm_quotes(option))
        if idx < 0:
            return False
        _rebuild(p, [(text[:idx], {}), (text[idx:idx + len(option)], {"underline": True, "bold": True}),
                     (text[idx + len(option):], {})])
        return True

    blanks = list(_BLANK.finditer(text))
    if blanks:
        # Varias líneas en blanco en el mismo párrafo: respuestas separadas por " | "
        parts = [v.strip() for v in value.split(" | ")] if len(blanks) > 1 else [value]
        segments, pos = [], 0
        for i, m in enumerate(blanks):
            segments.append((text[pos:m.start()], {}))
            answer = parts[i] if i < len(parts) else ""
            if answer:
                # Evitar que la respuesta quede pegada a la palabra de al lado ("_____video" → "her video")
                if m.start() > 0 and text[m.start() - 1].isalnum():
                    answer = " " + answer
                if m.end() < len(text) and text[m.end()].isalnum():
                    answer = answer + " "
            segments.append((answer or m.group(0), {"underline": bool(answer.strip())}))
            pos = m.end()
        segments.append((text[pos:], {}))
        _rebuild(p, segments)
        return True
    return False


def _insert_after(p: Paragraph, value: str) -> Paragraph:
    new_p = copy.deepcopy(p._p)
    for child in list(new_p):
        if child.tag.rsplit("}", 1)[-1] != "pPr":
            new_p.remove(child)
    p._p.addnext(new_p)
    para = Paragraph(new_p, p._parent)
    for i, chunk in enumerate(value.split("\n")):
        if i:
            para.add_run().add_break()
        para.add_run(chunk)
    return para


def fill(original: Path, answers_text: str, out_path: Path) -> tuple[Path, list[str]]:
    """Escribe las respuestas en una copia del Word. Devuelve (ruta, avisos de respuestas que no se pudieron colocar)."""
    doc = Document(str(original))
    paragraphs = list(doc.paragraphs)
    warnings = []
    last_inserted: dict[int, Paragraph] = {}

    for a in parse_answers(answers_text):
        value = a["value"]
        if "p" in a:
            if a["p"] >= len(paragraphs):
                warnings.append(f"[P{a['p']}] no existe en el documento")
                continue
            p = paragraphs[a["p"]]
            if a["after"] or not _fill_paragraph(p, value):
                if _UNDERLINE.match(value):
                    warnings.append(f"[P{a['p']}] no se encontró la opción a subrayar")
                    continue
                anchor = last_inserted.get(a["p"], p)
                last_inserted[a["p"]] = _insert_after(anchor, value)
        else:
            try:
                cell = doc.tables[a["t"]].cell(a["r"], a["c"])
            except IndexError:
                warnings.append(f"[T{a['t']} R{a['r']} C{a['c']}] no existe en el documento")
                continue
            target = cell.paragraphs[-1]
            if cell.text.strip() and not _BLANK.search(cell.text):
                target = _insert_after(target, value)
            elif not _fill_paragraph(target, value):
                target.add_run(value)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path, warnings
